#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the all-English report BaoCao_QuanLyThuVien_v3.2_en.docx from v3.2 by
translating the Vietnamese parts to English and swapping in the English diagrams."""
import os
from docx import Document
from docx.oxml.ns import qn

SRC = "BaoCao_QuanLyThuVien_v3.2.docx"
DST = "BaoCao_QuanLyThuVien_v3.2_en.docx"
HERE = os.path.dirname(os.path.abspath(__file__))

# ---- Vietnamese -> English (whole-paragraph / whole-cell, matched on stripped text) ----
T = {
 # captions
 "Hình 11.1 — Sơ đồ Use Case của hệ thống.": "Figure 11.1 — Use Case diagram of the system.",
 "Hình 11.2 — Sơ đồ lớp (mô hình miền dữ liệu).": "Figure 11.2 — Class diagram (domain model).",
 "Hình 11.3 — Sơ đồ hoạt động: luồng mượn sách tự phục vụ.": "Figure 11.3 — Activity diagram: the self-service borrowing flow.",
 "Hình 11.4 — Sơ đồ trạng thái: vòng đời đơn mượn (Loan).": "Figure 11.4 — State machine diagram: the loan lifecycle.",
 "Walkthrough (diễn giải):": "Walkthrough:",
 # 10.1 walkthrough
 "Bước 1 — Xuất trình thẻ: Bạn đọc đến phòng mượn xuất trình thẻ; CBTV kiểm tra tính hợp lệ của thẻ trước khi cho mượn tài liệu.":
   "Step 1 — Present card: The reader comes to the borrowing room and presents their card; library staff check that the card is valid before lending any document.",
 "Bước 2 — Tra cứu tài liệu: Bạn đọc tra tìm trên máy tra cứu tại thư viện hoặc tại elib.phenikaa-uni.edu.vn / library.phenikaa-uni.edu.vn (tìm theo nhan đề, tác giả, từ khoá, ISBN, ký hiệu phân loại, mã học phần; có tìm kiếm nâng cao và tra cứu qua giao thức Z39.50/SRU), rồi xem chi tiết biên mục (số ĐKCB, phân loại, thông tin xuất bản…).":
   "Step 2 — Search the catalogue: The reader searches on a library lookup terminal or at elib.phenikaa-uni.edu.vn / library.phenikaa-uni.edu.vn (by title, author, keyword, ISBN, classification code or course code; advanced search and Z39.50/SRU lookup are available), then views the catalogue record in detail (accession number, classification, publication data, etc.).",
 "Bước 3 — Viết phiếu yêu cầu: Điền phiếu Đề nghị cung cấp tài liệu BM.01.QT.03.TV.DV và mang đến bàn thủ thư (P104 – Nhà A4), hoặc ghi mượn trực tiếp trên phần mềm Koha.":
   "Step 3 — Fill in the request slip: Complete the document request slip BM.01.QT.03.TV.DV and bring it to the librarian's desk (Room P104 – Building A4), or record the loan directly in the Koha software.",
 "Bước 4 — Vào kho tìm tài liệu: CBTV vào kho chọn tài liệu theo phiếu yêu cầu của bạn đọc (kho được sắp xếp theo HD.01.TV.NV.04).":
   "Step 4 — Retrieve from the stacks: Library staff go into the stacks and pick the documents listed on the reader's request slip (the stacks are arranged per HD.01.TV.NV.04).",
 "Bước 5 — Làm thủ tục cho mượn: CBTV nhập thông tin vào module mượn của phần mềm / ghi sổ mượn–trả (BM.02.QT.03.TV.DV); bạn đọc kiểm tra tài liệu, nhận lại thẻ và ký nhận. Hạn mức và thời hạn mượn áp theo từng đối tượng (sinh viên / giảng viên / nghiên cứu viên) như quy định ở §3.2.":
   "Step 5 — Complete the loan: Library staff enter the data into the software's lending module / the borrow–return ledger (BM.02.QT.03.TV.DV); the reader checks the documents, takes the card back and signs. Quotas and loan durations apply per reader type (student / lecturer / researcher) as specified in §3.2.",
 "Bước 6 — Trả tài liệu: Bạn đọc trả tại P104 nhà A4. Gia hạn: mang tài liệu và thẻ đến quầy; chỉ được gia hạn nếu tài liệu còn ở trạng thái “sẵn sàng” và chưa có lượt đặt trước. Quá hạn: CBTV lập danh sách quá hạn, gửi thông báo về đơn vị (BM.04 & BM.05.QT.03.TV.DV).":
   "Step 6 — Return documents: The reader returns documents at Room P104, Building A4. Renewal: bring the document and card to the counter; renewal is only allowed if the document is still 'available' and has no pending reservation. Overdue: library staff compile an overdue list and send a notice to the reader's unit (BM.04 & BM.05.QT.03.TV.DV).",
 "Bước 7 — Làm thủ tục thu hồi: CBTV kiểm tra tài liệu khi trả, thao tác trả trên phần mềm / cho ký sổ; nếu hư hỏng, xử lý theo QT.05.TV.DV.":
   "Step 7 — Process the return: Library staff inspect the document on return, record the return in the software / have the reader sign the ledger; if it is damaged, handle it under QT.05.TV.DV.",
 "Bước 8 — Xếp tài liệu lên giá: Thủ thư xếp tài liệu trả lại đúng vị trí trên kệ theo HD.01.TV.DV.04.":
   "Step 8 — Reshelve documents: The librarian returns the document to its correct place on the shelf per HD.01.TV.DV.04.",
 # 10.2 walkthrough
 "Bước 1 — Đăng nhập: Xác thực bằng mã số + mật khẩu BCrypt (Spring Security). Tại “/”, hệ thống định tuyến: Admin → quản trị tài khoản/quyền; Thủ thư → dashboard nghiệp vụ; Bạn đọc (SV/GV/NCS) → trang tra cứu. Mọi URL bị chặn 2 lớp: cổng vai trò (SecurityConfig) và quyền chi tiết (PermissionService).":
   "Step 1 — Log in: Authentication uses the student/staff code + a BCrypt password (Spring Security). At “/”, the system routes: Admin → account/permission management; Librarian → the operations dashboard; Reader (student/lecturer/researcher) → the catalogue page. Every URL is guarded by two layers: the role gate (SecurityConfig) and granular permissions (PermissionService).",
 "Bước 2 — Tra cứu & xem chi tiết: Bạn đọc tìm theo từ khoá / loại tài liệu / chỉ còn sẵn; danh mục chỉ hiển thị những loại tài liệu mà vai trò được phép xem (cấu hình ở trang quy tắc của thủ thư).":
   "Step 2 — Search & view details: The reader searches by keyword / document type / availability; the catalogue shows only the document types the role is allowed to see (configured on the librarian's rules page).",
 "Bước 3 — Thêm vào giỏ: Giỏ lưu trong HTTP session, qua 6 lớp kiểm tra trước khi thêm — có quyền BORROW, chưa trùng trong giỏ, tổng ≤ 5 cuốn, không phải tài liệu nội sinh (RESTRICTED), loại tài liệu được phép xem, và sách còn bản sẵn.":
   "Step 3 — Add to cart: The cart lives in the HTTP session and passes six checks before an item is added — the BORROW permission, no duplicate in the cart, a total of ≤ 5 items, not an internal (RESTRICTED) document, the document type is visible to the role, and the book still has available copies.",
 "Bước 4 — Gửi phiếu mượn: RuleEngine kiểm tra 3 tầng (thẻ không bị khoá · không đang có tài liệu quá hạn · không vượt hạn mức theo từng loại). Hợp lệ thì tạo ngay đơn AWAITING_PICKUP, sinh mã lấy sách 6 ký tự, đặt hạn lấy 24 giờ và giữ chỗ (giảm availableCopies) — không cần thủ thư duyệt.":
   "Step 4 — Submit the loan request: The RuleEngine validates in three tiers (card not locked · no outstanding overdue document · within the per-type quota). If valid, an AWAITING_PICKUP loan is created immediately, a 6-character pickup code is generated, a 24-hour pickup deadline is set and the copies are held (availableCopies decremented) — with no librarian approval.",
 "Bước 5 — Chuẩn bị sách: Thủ thư xem danh sách đơn chờ lấy (kèm mã và hạn 24h) để soạn sách; có thể huỷ đơn kèm lý do (sẽ gửi cho bạn đọc và hoàn lại số lượng).":
   "Step 5 — Prepare the books: The librarian views the list of awaiting-pickup orders (with codes and 24-hour deadlines) to gather the books; they may cancel an order with a reason (which is sent to the reader and restores the held copies).",
 "Bước 6 — Tự xác nhận nhận sách: Tại “Đơn của tôi”, bạn đọc nhập lại mã lấy sách. Hệ thống kiểm tra đúng chủ đơn, đơn đang AWAITING_PICKUP và mã khớp (không phân biệt hoa/thường) → chuyển BORROWED và đặt hạn trả tính tại thời điểm lấy theo vai trò × loại tài liệu.":
   "Step 6 — Self-confirm pickup: On 'My Loans', the reader re-enters the pickup code. The system checks the correct owner, the AWAITING_PICKUP status and a matching code (case-insensitive) → the loan becomes BORROWED and the due date is computed at pickup time per role × document type.",
 "Bước 7 — Gia hạn: Nếu còn lượt (renewalCount < maxRenewals và maxRenewals > 0), hạn trả được cộng thêm renewalDays.":
   "Step 7 — Renew: If renewals remain (renewalCount < maxRenewals and maxRenewals > 0), the due date is extended by renewalDays.",
 "Bước 8 — Trả sách & thu phí: Thủ thư nhận trả từng cuốn, ghi tình trạng (GOOD/DAMAGED), tính phí trễ = số ngày quá hạn × 2.000đ (ChronoUnit.DAYS) và tăng lại số lượng. Khi mọi cuốn đã trả và mọi khoản phí đã thu (finePaid), đơn chuyển COMPLETED.":
   "Step 8 — Return & collect fines: The librarian processes each item's return, records its condition (GOOD/DAMAGED), computes the late fine = overdue days × 2,000 VND (ChronoUnit.DAYS) and restores the copy. Once every item is returned and every fine is paid (finePaid), the loan becomes COMPLETED.",
 "Bước 9 — Tác vụ nền: PickupExpiryScheduler chạy mỗi 10 phút huỷ các đơn quá 24h chưa lấy và hoàn số lượng; sendDueReminders() nhắc tài liệu sắp/đã quá hạn (đã hiện thực, chưa lên lịch).":
   "Step 9 — Background jobs: PickupExpiryScheduler runs every 10 minutes to cancel orders not collected within 24 hours and restore the copies; sendDueReminders() reminds readers of due-soon / overdue documents (implemented, not yet scheduled).",
 "Nhánh rẽ & luồng phụ trợ — Bạn đọc tự huỷ (chỉ khi chưa lấy), thủ thư huỷ (kèm lý do), hoặc scheduler tự huỷ; tất cả đều hoàn số lượng về kho. Tài liệu nội sinh (RESTRICTED) không mượn về — thủ thư lập phiếu Đọc tại chỗ (hạn trong ngày). Mọi vai trò: tự đổi mật khẩu và xem/đánh dấu thông báo in-app.":
   "Branches & supporting flows — A reader self-cancels (only before pickup), the librarian cancels (with a reason), or the scheduler auto-cancels; all restore the copies to stock. Internal (RESTRICTED) documents cannot be taken away — the librarian issues an in-house reading slip (due the same day). All roles: change their own password and view/mark in-app notifications.",
 # scattered VN inside English paragraphs
 "Each loan request generates a unique pickup code (mã mượn). Librarians simply prepare the books for that code.":
   "Each loan request generates a unique pickup code. Librarians simply prepare the books for that code.",
 "On 'My Loans', re-enter the pickup code and click 'Tôi đã nhận sách' → the loan becomes Đang mượn.":
   "On 'My Loans', re-enter the pickup code and click the 'I have received the books' button → the loan becomes Borrowed.",
 "Service hours: Monday–Friday, 08:00–11:30 and 13:30–16:30. The counter-based borrowing procedure has eight steps. The flowchart below reproduces the official regulation; the responsible party and the associated forms (biểu mẫu) are listed per step.":
   "Service hours: Monday–Friday, 08:00–11:30 and 13:30–16:30. The counter-based borrowing procedure has eight steps. The flowchart below reproduces the official regulation; the responsible party and the associated forms are listed per step.",
 "This section documents the document-borrowing workflow twice: first the original counter-based process of the Phenikaa University Library (regulation QT.03.TV.DV), then the digital self-service process implemented by our system, followed by a concise side-by-side comparison. CBTV = library staff (cán bộ thư viện).":
   "This section documents the document-borrowing workflow twice: first the original counter-based process of the Phenikaa University Library (regulation QT.03.TV.DV), then the digital self-service process implemented by our system, followed by a concise side-by-side comparison.",
 "Actors are grouped exactly as the RBAC design dictates — Reader (Student/Lecturer/Researcher), Librarian, Administrator — plus a non-human Scheduler actor for the background jobs. There is no role inheritance: each actor sees only its own use cases. Đăng nhập and Đổi mật khẩu apply to every authenticated role.":
   "Actors are grouped exactly as the RBAC design dictates — Reader (Student/Lecturer/Researcher), Librarian, Administrator — plus a non-human Scheduler actor for the background jobs. There is no role inheritance: each actor sees only its own use cases. Log in and Change password apply to every authenticated role.",
 # Table 7 messages (in-app, Vietnamese) -> English
 "Message (Vietnamese)": "Message (shown in Vietnamese in-app)",
 "Tài khoản của bạn đang bị khoá. Vui lòng liên hệ thủ thư.": "Your account is locked. Please contact the librarian.",
 "Bạn đang có N tài liệu quá hạn. Vui lòng trả sách trước khi mượn thêm.": "You have N overdue documents. Please return them before borrowing more.",
 "Giỏ sách trống. Vui lòng thêm sách trước khi gửi phiếu.": "Your cart is empty. Please add books before submitting.",
 "Tổng số tài liệu trong 1 lần mượn không được vượt quá 5 cuốn...": "A single loan may not exceed 5 documents in total...",
 "Tài liệu nội sinh (tem đỏ) chỉ được đọc tại chỗ, không được mượn về.": "Internal documents (red sticker) are read-on-site only and cannot be taken away.",
 "Bạn không có quyền mượn tài liệu loại: [tên loại].": "You are not allowed to borrow documents of type: [type name].",
 "Vượt quá giới hạn cho loại '[type]': tối đa N cuốn, bạn đang chọn M cuốn.": "Over the limit for type '[type]': max N, you selected M.",
 # ---- Table 25 (Phenikaa flowchart) ----
 "STT": "No.", "Lưu đồ (bước)": "Flowchart (step)", "Thực hiện": "Performed by",
 "Phối hợp": "Coordinates with", "Biểu mẫu": "Form",
 "Xuất trình thẻ": "Present card", "Bạn đọc": "Reader", "CBTV": "Library staff",
 "Tra cứu tài liệu": "Search the catalogue",
 "Viết phiếu yêu cầu": "Fill in the request slip",
 "BM.01.QT.03.TV.DV (hoặc ghi mượn trực tiếp trên PM Koha)": "BM.01.QT.03.TV.DV (or record the loan directly in Koha)",
 "Vào kho tìm tài liệu": "Retrieve from the stacks",
 "Làm thủ tục cho mượn": "Complete the loan",
 "BM.01.QT.03.TV.DV; Phần mềm Koha": "BM.01.QT.03.TV.DV; Koha software",
 "Trả tài liệu tại bàn CBTV": "Return at the staff desk",
 "BM.02.QT.03.TV.DV; Phần mềm Koha": "BM.02.QT.03.TV.DV; Koha software",
 "Làm thủ tục thu hồi / gia hạn / xử lý vi phạm (nếu có)": "Process return / renewal / handle violations (if any)",
 "Xếp TL lên giá / xử lý tài liệu theo tình trạng vi phạm": "Reshelve / handle documents per violation status",
 # ---- Table 26 (our flowchart) ----
 "Bước (lưu đồ)": "Step (flowchart)", "Hệ thống / Phối hợp": "System / Coordination",
 "Route / Tạo phẩm": "Route / Artifact",
 "Đăng nhập": "Log in", "Spring Security (BCrypt); định tuyến theo vai trò": "Spring Security (BCrypt); role-based routing",
 "Tra cứu & xem chi tiết": "Search & view details",
 "Chỉ hiển thị loại TL được phép xem (visibleDocTypes)": "Shows only visible document types (visibleDocTypes)",
 "Thêm vào giỏ": "Add to cart",
 "6 lớp kiểm tra: quyền · trùng · ≤5 · không nội sinh · được phép · còn sách": "6 checks: permission · duplicate · ≤5 · not restricted · visible · in stock",
 "Gửi phiếu mượn (tự phục vụ)": "Submit loan request (self-service)",
 "RuleEngine validate 3 tầng → tạo đơn AWAITING_PICKUP + mã 6 ký tự + giữ chỗ (copies−1)": "RuleEngine 3-tier validate → create AWAITING_PICKUP loan + 6-char code + hold copies (copies−1)",
 "Chuẩn bị sách theo mã": "Prepare books by code", "Thủ thư": "Librarian",
 "Xem danh sách chờ lấy theo mã & hạn 24h": "View the awaiting-pickup list by code & 24h deadline",
 "Tự xác nhận nhận sách": "Self-confirm pickup",
 "confirmPickup: đúng chủ đơn + đúng mã → BORROWED, đặt hạn trả theo vai trò × loại TL": "confirmPickup: correct owner + correct code → BORROWED, due date set by role × doc type",
 "Gia hạn (nếu còn lượt)": "Renew (if allowed)",
 "canRenew: renewalCount < maxRenewals; +renewalDays": "canRenew: renewalCount < maxRenewals; +renewalDays",
 "Trả sách & thu phí": "Return & collect fines",
 "returnBook: phí = ngày trễ × 2.000đ, tăng copies; pay-fine; đơn → COMPLETED khi đã trả hết & thu hết phí": "returnBook: fine = overdue days × 2,000 VND, restore copies; pay-fine; loan → COMPLETED when all returned & all fines paid",
 "Tác vụ nền (tự động)": "Background jobs (automatic)", "Hệ thống": "System",
 "Mỗi 10 phút huỷ đơn quá 24h chưa lấy (hoàn copies); nhắc hạn trả sắp/đã quá hạn": "Every 10 min cancels orders not collected within 24h (restores copies); reminds of due-soon/overdue items",
 "Nhánh rẽ của đơn": "Order branches", "Bạn đọc / Thủ thư / Hệ thống": "Reader / Librarian / System",
 "Tự huỷ (khi chưa lấy) · thủ thư huỷ kèm lý do · scheduler tự huỷ — đều hoàn copies. Tài liệu nội sinh: lập phiếu Đọc tại chỗ": "Self-cancel (before pickup) · librarian cancel with reason · scheduler auto-cancel — all restore copies. Internal documents: issue an in-house reading slip",
 # ---- Table 27 (comparison) ----
 "Tiêu chí": "Criterion", "Quy trình Phenikaa (thủ công / Koha)": "Phenikaa process (manual / Koha)",
 "Hệ thống của chúng tôi (web tự phục vụ)": "Our system (self-service web)",
 "Khởi tạo mượn": "Loan initiation",
 "Bạn đọc viết phiếu / CBTV ghi mượn; cần thủ thư thao tác": "Reader fills a slip / staff record the loan; requires librarian action",
 "Tự phục vụ: bạn đọc tự tạo đơn, không cần duyệt": "Self-service: the reader creates the order, no approval needed",
 "Định danh đơn": "Order identity", "Phiếu giấy / sổ mượn–trả": "Paper slip / borrow–return ledger",
 "Mã lấy sách 6 ký tự sinh tự động": "Auto-generated 6-character pickup code",
 "Xác nhận giao sách": "Hand-over confirmation", "Thủ thư trao tay tại quầy": "Librarian hands over at the counter",
 "Bạn đọc tự nhập mã để xác nhận (confirmPickup)": "Reader self-confirms by entering the code (confirmPickup)",
 "Giữ chỗ tồn kho": "Stock holding", "Thủ công khi ghi mượn": "Manual when recording the loan",
 "Tự động trừ availableCopies ngay khi tạo đơn": "availableCopies decremented automatically when the order is created",
 "Hạn mức & hạn trả": "Quota & due date", "Quy định văn bản; CBTV áp dụng thủ công": "Written regulation; staff apply manually",
 "RuleEngine tự áp theo vai trò × loại TL (cấu hình trên UI)": "RuleEngine applies them automatically by role × doc type (configured in the UI)",
 "Quá hạn lấy sách": "Pickup expiry", "Không có khái niệm": "No such concept",
 "Scheduler tự huỷ sau 24h và hoàn sách": "Scheduler auto-cancels after 24h and restores the books",
 "Phí trễ hạn": "Late fine", "Tính thủ công": "Calculated manually",
 "Tự động: ngày trễ × 2.000đ; theo dõi đã thu (finePaid)": "Automatic: overdue days × 2,000 VND; payment tracked (finePaid)",
 "Gia hạn": "Renewal", "Tại quầy; kiểm tra “sẵn sàng” / đặt trước": "At the counter; checks 'available' / reservations",
 "Online nếu còn lượt (canRenew)": "Online if renewals remain (canRenew)",
 "Nhắc quá hạn trả": "Overdue reminders", "CBTV lập danh sách, gửi công văn (BM.04/05)": "Staff compile a list and send an official notice (BM.04/05)",
 "Thông báo in-app theo sự kiện; sendDueReminders (chưa lên lịch)": "Event-driven in-app notifications; sendDueReminders (not yet scheduled)",
 "Tài liệu nội sinh (dấu đỏ)": "Internal documents (red mark)", "Không cho mượn về": "Cannot be taken away",
 "RESTRICTED: chặn mượn; thủ thư lập phiếu Đọc tại chỗ": "RESTRICTED: borrowing blocked; librarian issues an in-house reading slip",
 "Phân quyền": "Authorisation", "Theo vai trò CBTV": "By staff role",
 "RBAC 2 lớp: cổng vai trò + ma trận quyền chi tiết": "Two-layer RBAC: role gate + granular permission matrix",
 "Vết & trạng thái": "Trace & status", "Sổ sách / email đơn vị": "Ledgers / unit emails",
 "Vòng đời đơn có trạng thái rõ ràng + thông báo in-app": "A clear loan lifecycle with explicit statuses + in-app notifications",
}

def apply_runs(p):
    full = p.text
    key = full.strip()
    if key in T and p.runs:
        lead = full[:len(full)-len(full.lstrip())]
        trail = full[len(full.rstrip()):]
        p.runs[0].text = lead + T[key] + trail
        for r in p.runs[1:]:
            r.text = ""

doc = Document(SRC)
for p in doc.paragraphs:
    apply_runs(p)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                apply_runs(p)

# ---- swap the 4 embedded images for the English renders ----
img_map = {  # caption-key (English now) -> EN png
 "Figure 11.1 — Use Case diagram of the system.": "uml_usecase_en.png",
 "Figure 11.2 — Class diagram (domain model).": "uml_class_en.png",
 "Figure 11.3 — Activity diagram: the self-service borrowing flow.": "uml_activity_en.png",
 "Figure 11.4 — State machine diagram: the loan lifecycle.": "uml_state_en.png",
}
ps = doc.paragraphs
for i, p in enumerate(ps):
    blips = p._p.findall('.//'+qn('a:blip'))
    if not blips:
        continue
    cap = ps[i+1].text.strip() if i+1 < len(ps) else ''
    png = img_map.get(cap)
    if not png:
        continue
    rid = blips[0].get(qn('r:embed'))
    part = doc.part.related_parts[rid]
    with open(os.path.join(HERE, png), 'rb') as f:
        part._blob = f.read()
    print(f"swapped image at P#{i} -> {png}")

doc.save(DST)
print("SAVED", DST)
